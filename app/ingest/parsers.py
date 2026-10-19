"""Читання PDF (з текстовим шаром), DOCX і TXT."""
import io
import re

import pymupdf as fitz
from docx import Document

MIN_PAGE_CHARS = 30  # менше — сторінка вважається сканом


def clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def parse_txt(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return clean(data.decode(enc))
        except UnicodeDecodeError:
            continue
    return clean(data.decode("utf-8", errors="replace"))


def parse_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return clean("\n".join(parts))


def parse_pdf(data: bytes) -> tuple[list[str], list[int]]:
    """Повертає текст кожної сторінки та номери сторінок без текстового шару (скани)."""
    pages: list[str] = []
    scans: list[int] = []
    with fitz.open(stream=data, filetype="pdf") as pdf:
        for i, page in enumerate(pdf):
            text = clean(page.get_text())
            pages.append(text)
            if len(text) < MIN_PAGE_CHARS:
                scans.append(i)
    return pages, scans


def render_page_jpeg(data: bytes, page_no: int, dpi: int = 200) -> bytes:
    with fitz.open(stream=data, filetype="pdf") as pdf:
        pix = pdf[page_no].get_pixmap(dpi=dpi)
        return pix.tobytes("jpeg", jpg_quality=88)
